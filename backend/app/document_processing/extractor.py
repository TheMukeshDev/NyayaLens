"""Text extraction from PDF, DOCX and image files.

Dispatch by validated MIME type:

* ``application/pdf`` — pypdf, per page.
* ``application/vnd.openxmlformats-officedocument.wordprocessingml.document`` —
  python-docx; page numbers are approximated from explicit page breaks (real
  pagination requires document rendering, which is out of scope).
* ``image/jpeg`` / ``image/png`` — the injected :class:`OcrEngine`.

Extraction is lossy by design. Failures on an individual page become
:class:`ExtractionUncertainty` records instead of corrupting the whole
document, and empty results are never padded with invented text.
"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any

from app.document_processing.models import (
    EXTRACTION_FAILED,
    NO_TEXT,
    ExtractedPage,
    ExtractedText,
    ExtractionUncertainty,
)
from app.document_processing.ocr import OcrEngine, OcrResult

logger = logging.getLogger("app.processing.extractor")

MIME_TO_FORMAT: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "image/jpeg": "image",
    "image/png": "image",
}

_PAGE_BREAK = "\x00PAGE\x00"  # sentinel for python-docx page-break runs


class ExtractionError(RuntimeError):
    """Raised when a file cannot be extracted at all (not per-page noise)."""


class TextExtractor:
    """Extracts structured page text from validated document bytes."""

    def __init__(self, *, ocr_engine: OcrEngine | None = None) -> None:
        self._ocr_engine = ocr_engine

    # -- public ----------------------------------------------------------------

    def extract(self, *, content: bytes, mime_type: str) -> ExtractedText:
        format_name = MIME_TO_FORMAT.get(mime_type or "")
        if format_name is None:
            raise ExtractionError(f"Unsupported content type: {mime_type or 'unknown'}")

        if format_name == "pdf":
            return self._extract_pdf(content)
        if format_name == "docx":
            return self._extract_docx(content)
        return self._extract_image(content)

    # -- pdf -------------------------------------------------------------------

    def _extract_pdf(self, content: bytes) -> ExtractedText:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - declared dependency
            raise ExtractionError("pypdf is not installed.") from exc

        try:
            reader = PdfReader(BytesIO(content), strict=False)
        except Exception as exc:  # unreadable/malformed PDF
            logger.error("PDF open failed", exc_info=True)
            raise ExtractionError(f"Could not read PDF: {type(exc).__name__}.") from exc

        MAX_PDF_PAGES = 500
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractionError(f"PDF exceeds maximum allowed length of {MAX_PDF_PAGES} pages.")

        pages: list[ExtractedPage] = []
        for page_index, page in enumerate(reader.pages, start=1):
            number = reader.get_page_number(page) or page_index
            page_uncertainties: list[ExtractionUncertainty] = []
            try:
                text = page.extract_text() or ""
            except Exception:
                logger.warning("PDF page %s text extraction failed", number, exc_info=True)
                text = ""
                page_uncertainties.append(
                    ExtractionUncertainty(
                        page=number,
                        code=EXTRACTION_FAILED,
                        message=f"Text extraction failed on page {number}.",
                    )
                )
            if not text.strip():
                page_uncertainties.append(
                    ExtractionUncertainty(
                        page=number,
                        code=NO_TEXT,
                        message=f"No text was extracted from page {number}.",
                    )
                )
            pages.append(
                ExtractedPage(
                    page_number=number,
                    text=text,
                    source="text",
                    uncertainties=page_uncertainties,
                )
            )
        return ExtractedText(format="pdf", pages=pages or [self._empty_page(1)])

    # -- docx ------------------------------------------------------------------

    def _extract_docx(self, content: bytes) -> ExtractedText:
        try:
            import docx as docx_lib
        except ImportError as exc:  # pragma: no cover - declared dependency
            raise ExtractionError("python-docx is not installed.") from exc

        import zipfile
        try:
            with zipfile.ZipFile(BytesIO(content)) as zf:
                if len(zf.infolist()) > 2000 or sum(z.file_size for z in zf.infolist()) > 60 * 1024 * 1024:
                    raise ExtractionError("DOCX archive exceeds safe expansion limits (potential decompression bomb).")
        except zipfile.BadZipFile as exc:
            raise ExtractionError(f"Could not read DOCX: {type(exc).__name__}.") from exc

        try:
            document = docx_lib.Document(BytesIO(content))
        except Exception as exc:
            logger.error("DOCX open failed", exc_info=True)
            raise ExtractionError(f"Could not read DOCX: {type(exc).__name__}.") from exc

        pages: list[list[str]] = [[]]
        for paragraph in list(document.paragraphs) + self._table_paragraphs(document):
            segments = _split_page_breaks(_paragraph_text(paragraph._p))
            for index, segment in enumerate(segments):
                if index > 0:
                    pages.append([])
                # Empty paragraphs are preserved as blank lines; structural
                # parsing (section/clause detection) relies on them.
                pages[-1].append(segment)

        extracted: list[ExtractedPage] = []
        for number, page_lines in enumerate(pages, start=1):
            text = "\n".join(page_lines).strip()
            uncertainties = (
                [
                    ExtractionUncertainty(
                        page=number,
                        code=NO_TEXT,
                        message=f"No text was extracted on page {number}.",
                    )
                ]
                if not text
                else []
            )
            extracted.append(
                ExtractedPage(
                    page_number=number,
                    text=text,
                    source="text",
                    uncertainties=uncertainties,
                )
            )
        return ExtractedText(format="docx", pages=extracted)

    @staticmethod
    def _table_paragraphs(document: Any) -> list[Any]:
        """Return every paragraph inside tables (python-docx ``_Paragraph``)."""
        paragraphs: list[Any] = []
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        return paragraphs

    # -- images ----------------------------------------------------------------

    def _extract_image(self, content: bytes) -> ExtractedText:
        if self._ocr_engine is None:
            raise ExtractionError(
                "No OCR engine is configured; scanned images cannot be processed."
            )
        result: OcrResult = self._ocr_engine.ocr(content)

        uncertainties: list[ExtractionUncertainty] = []
        if result.confidence is not None and result.confidence < 0.5:
            uncertainties.append(
                ExtractionUncertainty(
                    page=1,
                    code="LOW_OCR_CONFIDENCE",
                    message=(
                        f"OCR confidence is low "
                        f"({result.confidence:.2f}) on page 1; text may be unreliable."
                    ),
                )
            )
        page = ExtractedPage(
            page_number=1,
            text=result.text,
            source="ocr",
            confidence=result.confidence,
            uncertainties=uncertainties,
        )
        return ExtractedText(format="image", pages=[page], ocr_used=True)

    # -- helpers ---------------------------------------------------------------

    @staticmethod
    def _empty_page(page_number: int) -> ExtractedPage:
        return ExtractedPage(
            page_number=page_number,
            text="",
            source="text",
            uncertainties=[
                ExtractionUncertainty(
                    page=page_number, code=NO_TEXT, message="The document is empty."
                )
            ],
        )


def _paragraph_text(p_element: Any) -> str:
    """Return paragraph text with explicit page breaks as sentinel values."""
    from docx.oxml.ns import qn

    parts: list[str] = []
    for node in p_element.iter():
        tag = node.tag
        if tag == qn("w:t"):
            parts.append(getattr(node, "text", None) or "")
        elif tag == qn("w:br"):
            if node.get(qn("w:type")) == "page":
                parts.append(_PAGE_BREAK)
            else:
                parts.append("\n")
    return "".join(parts)


def _split_page_breaks(paragraph_text: str) -> list[str]:
    """Split on page-break sentinels, dropping empty sentinel-only parts."""
    raw = paragraph_text.split(_PAGE_BREAK)
    segments: list[str] = []
    for segment in raw:
        if segment:
            segments.append(segment)
    # A trailing page break still ends the paragraph on a new page.
    if raw and paragraph_text.rstrip().endswith(_PAGE_BREAK):
        segments.append("")
    return segments or [""]


def aggregate_uncertainties(text: ExtractedText) -> list[ExtractionUncertainty]:
    """Flatten page-level uncertainties into a single top-level list."""
    uncertainties = list(text.uncertainties)
    for page in text.pages:
        uncertainties.extend(page.uncertainties)
    return uncertainties