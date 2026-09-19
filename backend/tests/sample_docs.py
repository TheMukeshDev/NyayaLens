"""Synthetic document builders for hermetic processing tests.

Produces real, readable bytes (not mocks): PDFs with literal text content
streams and correct xref tables, DOCX files via python-docx, and PNG images
via Pillow. The fake OCR engine used by tests lives in ``test_processing.py``.
"""

from __future__ import annotations

import io

import docx
from docx.enum.text import WD_BREAK
from PIL import Image, ImageDraw

# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def pdf_bytes(*pages: str) -> bytes:
    """Return a valid one-page-per-*pages* PDF carrying literal text.

    Page text lines become separate ``Tj`` operators so pypdf extracts them as
    line-separated text. xref offsets are computed as objects are written, so
    the emitted file is fully readable by pypdf.
    """
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}

    def add_object(object_number: int, body: str) -> None:
        payload = (
            f"{object_number} 0 obj\n{body}\nendobj\n".encode("latin-1")
        )
        offsets[object_number] = out.tell()
        out.write(payload)

    add_object(1, "<< /Type /Catalog /Pages 2 0 R >>")
    font_id = 3 + len(pages) * 2

    kids: list[str] = []
    for page_number, page_text in enumerate(pages):
        page_id = 3 + page_number * 2
        content_id = page_id + 1
        kids.append(f"{page_id} 0 R")
        stream = _pdf_content_stream(page_text).encode("latin-1")
        add_object(
            page_id,
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>",
        )
        add_object(
            content_id,
            f"<< /Length {len(stream)} >>\nstream\n"
            + stream.decode("latin-1")
            + "\nendstream",
        )

    add_object(font_id, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    add_object(
        2, f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    )

    xref_offset = out.tell()
    out.write(f"xref\n0 {len(offsets) + 1}\n".encode("latin-1"))
    out.write(b"0000000000 65535 f \n")
    for object_number in sorted(offsets):
        out.write(f"{offsets[object_number]:010d} 00000 n \n".encode("latin-1"))
    out.write(
        (
            f"trailer\n<< /Size {len(offsets) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("latin-1")
    )
    return out.getvalue()


def _pdf_content_stream(text: str) -> str:
    """Emit one operator per line, preserving blank lines as empty lines.

    Blank lines are semantically meaningful to the pipeline's section/clause
    detection; pypdf ignores empty ``Tj`` strings, so blank lines are encoded
    as an explicit ``Td`` line move.
    """
    body = ["BT /F1 12 Tf 50 730 Td 14 TL"]
    for line in text.split("\n"):
        if line.strip():
            body.append(f"({_escape_pdf_text(line)}) Tj T*")
        else:
            body.append("0 -14 Td")
    body.append("ET")
    return "\n".join(body)


def _escape_pdf_text(value: str) -> str:
    escaped: list[str] = []
    for char in value:
        codepoint = ord(char)
        if codepoint > 127:
            escaped.append(f"\\{codepoint:o}")
        elif char == "\\":
            escaped.append("\\\\")
        elif char == "(":
            escaped.append("\\(")
        elif char == ")":
            escaped.append("\\)")
        else:
            escaped.append(char)
    return "".join(escaped)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def docx_bytes(*paragraphs: str, insert_page_break_after: list[int] | None = None) -> bytes:
    """Return a DOCX whose paragraphs are *paragraphs* ('' = blank line).

    ``insert_page_break_after`` lists paragraph indexes after which an explicit
    page break is inserted (page 2 detection relies on these).
    """
    document = docx.Document()
    for index, paragraph_text in enumerate(paragraphs):
        paragraph = document.add_paragraph(paragraph_text)
        if insert_page_break_after and index in insert_page_break_after:
            paragraph.add_run().add_break(WD_BREAK.PAGE)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PNG
# ---------------------------------------------------------------------------


def png_bytes(*lines: str, width: int = 1200, height: int = 800) -> bytes:
    """Return a PNG image with *lines* drawn as white text on black."""
    image = Image.new("RGB", (width, height), "black")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        draw.text((60, 60 + index * 70), line, fill="white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()